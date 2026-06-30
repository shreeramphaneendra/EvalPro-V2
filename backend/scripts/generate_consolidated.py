#!/usr/bin/env python3
import sys, json, io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def romanize(n):
    m={1:'I',2:'II',3:'III',4:'IV',5:'V',6:'VI',7:'VII',8:'VIII'}
    return m.get(int(n),str(n))

def thin():
    s=Side(style='thin')
    return Border(left=s,right=s,top=s,bottom=s)

def fill(hex_c):
    return PatternFill('solid',fgColor=hex_c)

def set_cell(ws, row, col, value, font=None, align=None, f=None, border=None):
    c=ws.cell(row=row,column=col)
    c.value=value
    if font:   c.font=font
    if align:  c.alignment=align
    if f:      c.fill=f
    if border: c.border=border
    return c

def make_sheet(wb, sec_data, dept, year, program, semester, all_subs, theory_subs, elective_subs, lab_subs):
    label = sec_data['label']
    students = sec_data['students']
    section_num = str(sec_data['sectionNum'])

    ws = wb.create_sheet(title=label[:31])

    NC = len(all_subs)
    D  = 4          # first subject col (D=4)
    TC = D+NC       # Total col
    AC = D+NC+1     # Activity Points col
    LC = AC

    # col widths
    ws.column_dimensions['A'].width=6
    ws.column_dimensions['B'].width=18
    ws.column_dimensions['C'].width=32
    for i in range(NC): ws.column_dimensions[get_column_letter(D+i)].width=9
    ws.column_dimensions[get_column_letter(TC)].width=8
    ws.column_dimensions[get_column_letter(AC)].width=15

    LL = get_column_letter(LC)
    BF = Font(name='Calibri',bold=True,size=10)
    NF = Font(name='Calibri',size=10)
    CT = Alignment(horizontal='center',vertical='center',wrap_text=True)
    LT = Alignment(horizontal='left',vertical='center',wrap_text=True)

    # ROW 1 — college name
    ws.merge_cells(f'A1:{LL}1')
    set_cell(ws,1,1,'CHAITANYA BHARATHI INSTITUTE OF TECHNOLOGY(Autonomous), HYDERABAD-75',
             Font(name='Calibri',bold=True,size=13,color='FFFFFF'),CT,fill('1F4E79'))
    ws.row_dimensions[1].height=22

    # ROW 2 — title
    ws.merge_cells(f'A2:{LL}2')
    set_cell(ws,2,1,'CONSOLIDATED CIE MARKS',Font(name='Calibri',bold=True,size=12,color='FFFFFF'),CT,fill('2E75B6'))
    ws.row_dimensions[2].height=18

    # ROW 3 — program + dept
    ws.merge_cells('A3:D3')
    set_cell(ws,3,1,'Program:(B.E/B.Tech/MBA/MCA) ',BF,LT)
    ws.merge_cells('E3:G3')
    set_cell(ws,3,5,'Name of the Department:',BF,LT)
    ws.merge_cells(f'H3:{LL}3')
    set_cell(ws,3,8,dept,Font(name='Calibri',bold=True,size=10,color='C00000'),LT)
    ws.row_dimensions[3].height=15

    # ROW 4 — year + class + section
    ws.merge_cells('A4:D4')
    set_cell(ws,4,1,f'Academic Year: {year}',BF,LT)
    ws.merge_cells('E4:G4')
    set_cell(ws,4,5,'Class & Semester:',BF,LT)
    ws.merge_cells('H4:I4')
    set_cell(ws,4,8,f'B.E. & {romanize(semester)}',Font(name='Calibri',bold=True,size=10,color='C00000'),CT)
    set_cell(ws,4,10,'Section:',BF,LT)
    ws.merge_cells(f'K4:{LL}4')
    set_cell(ws,4,11,label,Font(name='Calibri',bold=True,size=10,color='C00000'),LT)
    ws.row_dimensions[4].height=15

    # Fills for subject groups
    TF=fill('D9E1F2')  # theory blue
    EF=fill('FFF2CC')  # elective yellow
    LF=fill('E2EFDA')  # lab green
    HF=fill('D9D9D9')  # gray header

    ts=len(theory_subs); es=len(elective_subs); ls_=len(lab_subs)

    def sub_fill(i):
        if i<ts: return TF
        if i<ts+es: return EF
        return LF

    # ROW 5 — group labels
    ws.row_dimensions[5].height=14
    if ts>0:
        a=get_column_letter(D); b=get_column_letter(D+ts-1)
        set_cell(ws,5,D,'Theory',BF,CT,TF)
        if ts>1: ws.merge_cells(f'{a}5:{b}5')
    if es>0:
        a=get_column_letter(D+ts)
        # row 5 blank for elective — label goes in row 6
        set_cell(ws,5,D+ts,'',None,None,EF)
        if es>1:
            b=get_column_letter(D+ts+es-1)
            ws.merge_cells(f'{a}5:{b}5')
    if ls_>0:
        a=get_column_letter(D+ts+es); b=get_column_letter(D+ts+es+ls_-1)
        set_cell(ws,5,D+ts+es,'Practicals',BF,CT,LF)
        if ls_>1: ws.merge_cells(f'{a}5:{b}5')

    # ROW 6 — elective label
    ws.row_dimensions[6].height=14
    for i in range(NC): set_cell(ws,6,D+i,'',None,None,sub_fill(i))
    if es>0:
        a=get_column_letter(D+ts); b=get_column_letter(D+ts+es-1)
        set_cell(ws,6,D+ts,'Professional Elective - I',Font(name='Calibri',bold=True,size=9),CT,EF)
        if es>1: ws.merge_cells(f'{a}6:{b}6')

    # ROWS 7-9 — codes / names / maxmarks  (S.NO/Roll/Name merged 7-9)
    for r in [7,8,9]: ws.row_dimensions[r].height=15

    ws.merge_cells('A7:A9'); set_cell(ws,7,1,'S.NO',BF,CT,HF)
    ws.merge_cells('B7:B9'); set_cell(ws,7,2,'Roll No',BF,CT,HF)
    ws.merge_cells('C7:C9'); set_cell(ws,7,3,'Name of the Student',BF,CT,HF)
    ws.merge_cells(f'{get_column_letter(TC)}7:{get_column_letter(TC)}9')
    set_cell(ws,7,TC,'Total',BF,CT,fill('FCE4D6'))
    ws.merge_cells(f'{get_column_letter(AC)}7:{get_column_letter(AC)}9')
    set_cell(ws,7,AC,'Activity Points',BF,CT,HF)

    for i,sub in enumerate(all_subs):
        sf=sub_fill(i)
        set_cell(ws,7,D+i,sub['code'],Font(name='Calibri',bold=True,size=9),CT,sf)
        set_cell(ws,8,D+i,sub.get('shortName',sub['name'][:8]),Font(name='Calibri',bold=True,size=8),CT,sf)
        mx=50 if sub['type']=='lab' else 40
        set_cell(ws,9,D+i,mx,Font(name='Calibri',bold=True,size=9),CT,sf)

    grand_max=sum(50 if s['type']=='lab' else 40 for s in all_subs)
    set_cell(ws,9,TC,grand_max,Font(name='Calibri',bold=True,size=9),CT,fill('FCE4D6'))

    # Apply borders to header rows 7-9
    for r in range(7,10):
        for ci in range(1,LC+1):
            ws.cell(row=r,column=ci).border=thin()

    # DATA ROWS
    DR=10
    for idx,student in enumerate(students):
        r=DR+idx
        ws.row_dimensions[r].height=14
        even=idx%2==1; rf=fill('F2F2F2') if even else None

        set_cell(ws,r,1,idx+1,NF,CT,rf)
        set_cell(ws,r,2,student['usn'],Font(name='Calibri',size=9),CT,rf)
        set_cell(ws,r,3,student['name'].upper(),NF,LT,rf)

        for i,sub in enumerate(all_subs):
            score=student['marks'].get(sub['_id'])
            f_=rf if rf else None
            c=ws.cell(row=r,column=D+i)
            c.value=score if score is not None else None
            c.font=NF; c.alignment=CT
            if f_: c.fill=f_

        # Total formula
        fc=get_column_letter(D); lc_=get_column_letter(D+NC-1)
        c=ws.cell(row=r,column=TC)
        c.value=f'=SUM({fc}{r}:{lc_}{r})'
        c.font=Font(name='Calibri',bold=True,size=10); c.alignment=CT; c.fill=fill('FCE4D6')

        # Activity Points
        c=ws.cell(row=r,column=AC)
        c.value=student.get('activityPoints') or None
        c.font=NF; c.alignment=CT
        if rf: c.fill=rf

        # borders
        for ci in range(1,LC+1):
            ws.cell(row=r,column=ci).border=thin()

    last_data=DR+len(students)-1

    # FACULTY TABLE
    fr=last_data+2
    ws.row_dimensions[fr].height=14
    for ci,h in enumerate(['Sub-Code','Course Name','Name of the Faculty','Mail ID','Mobile Number'],1):
        set_cell(ws,fr,ci,h,BF,CT,HF,thin())

    r=fr+1
    for grp_label,subs in [('Core Subjects',theory_subs),('Professional Elective – I (PE-I)',elective_subs),('Practicals',lab_subs)]:
        if not subs: continue
        ws.merge_cells(f'A{r}:E{r}')
        set_cell(ws,r,1,grp_label,Font(name='Calibri',bold=True,size=10),None,fill('BDD7EE'))
        r+=1
        for sub in subs:
            teachers=sub.get('teachers',{}).get(section_num,[]) or sub.get('teachers',{}).get('all',[]) or [{'name':'—','email':'','phone':''}]
            for ti,t in enumerate(teachers):
                if ti==0:
                    set_cell(ws,r,1,sub['code'],NF,None,None,thin())
                    set_cell(ws,r,2,sub['name'],NF,None,None,thin())
                set_cell(ws,r,3,t.get('name','—'),NF,None,None,thin())
                set_cell(ws,r,4,t.get('email',''),NF,None,None,thin())
                set_cell(ws,r,5,str(t.get('phone','')) if t.get('phone') else '',NF,None,None,thin())
                r+=1

    ws.freeze_panes='D10'

def main():
    payload=json.loads(sys.stdin.buffer.read())
    dept=payload['dept']; year=payload['year']
    program=payload['program']; semester=payload['semester']
    all_subs=payload['subjects']
    theory_subs=[s for s in all_subs if s['type']=='theory']
    elective_subs=[s for s in all_subs if s['type']=='elective']
    lab_subs=[s for s in all_subs if s['type']=='lab']

    wb=Workbook(); wb.remove(wb.active)
    for sec in payload['sections']:
        make_sheet(wb,sec,dept,year,program,semester,all_subs,theory_subs,elective_subs,lab_subs)
    buf=io.BytesIO(); wb.save(buf)
    sys.stdout.buffer.write(buf.getvalue())

if __name__=='__main__': main()
